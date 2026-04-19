"""
学期服务
"""

import json
import io
import logging
from datetime import date, timedelta
from django.db import transaction, models
from django.core.paginator import Paginator
from django.core.cache import cache
from django.core.serializers import serialize
from django.http import HttpResponse
from django.utils import timezone
from django.utils.encoding import escape_uri_path
from apps.core.exceptions import ValidationError, NotFoundError, PermissionDenied
from apps.schedules.models import Semester, TermArchive
from apps.laboratories.models import Laboratory, Equipment
from apps.records.models import UsageRecord
from apps.maintenance.models import WorkOrder
from apps.users.models import User, SystemSetting
from apps.schedules.models import Schedule
from common.services.cache_service import CacheKeyManager, CacheInvalidator
from common.decorators import cached_method

logger = logging.getLogger(__name__)


class SemesterService:
    """学期服务"""

    @cached_method(timeout=60, key_prefix='semester:list')
    def get_semester_list(
        self,
        requester,
        page: int = 1,
        page_size: int = 20,
        no_page: bool = False
    ) -> dict:
        queryset = Semester.objects.filter(is_deleted=False)
        queryset = queryset.order_by('-start_date')
        
        archived_terms = []
        archived_semester_ids = TermArchive.objects.values_list('semester_id', flat=True).distinct()
        archived_semesters = Semester.objects.filter(id__in=archived_semester_ids).order_by('-end_date')
        for sem in archived_semesters:
            archived_terms.append({
                'id': sem.id,
                'name': sem.name,
                'start_date': sem.start_date.strftime('%Y-%m-%d'),
                'end_date': sem.end_date.strftime('%Y-%m-%d'),
                'stats': self._get_archived_dashboard_stats(sem),
            })
        
        if no_page:
            return {
                'list': [self._format_semester(s) for s in queryset],
                'archived_terms': archived_terms,
                'pagination': {
                    'total': queryset.count(),
                }
            }
        
        paginator = Paginator(queryset, page_size)
        page_obj = paginator.get_page(page)
        
        return {
            'list': [self._format_semester(s) for s in page_obj],
            'archived_terms': archived_terms,
            'pagination': {
                'total': paginator.count,
                'page': page,
                'page_size': page_size,
                'total_pages': paginator.num_pages,
            }
        }

    def get_semester_detail(self, requester, semester_id: int) -> dict:
        try:
            semester = Semester.objects.get(id=semester_id, is_deleted=False)
        except Semester.DoesNotExist:
            raise NotFoundError('学期不存在')
        
        return self._format_semester_detail(semester)

    def get_current_semester(self, requester) -> dict:
        semester = Semester.get_current()
        if not semester:
            return {'current_semester': None}
        return {'current_semester': self._format_semester(semester)}

    def get_semester_options(self, requester) -> list:
        semesters = Semester.objects.filter(is_deleted=False).order_by('-start_date')
        options = [{'id': '', 'text': '全部学期'}]
        for s in semesters:
            label = s.name
            if s.is_current:
                label += ' (当前)'
            options.append({'id': s.id, 'text': label})
        return options

    @transaction.atomic
    def create_semester(self, requester, data: dict) -> Semester:
        if not requester.is_super_admin:
            raise PermissionDenied('无权限创建学期')
        
        name = data.get('name', '').strip()
        code = data.get('code', '').strip()
        
        if not name:
            raise ValidationError('学期名称不能为空')
        
        if not code:
            code = name.lower().replace(' ', '_').replace('-', '_')
        
        if Semester.objects.filter(code=code).exists():
            counter = 1
            original_code = code
            while Semester.objects.filter(code=code).exists():
                code = f"{original_code}_{counter}"
                counter += 1
        
        if Semester.objects.filter(name=name).exists():
            raise ValidationError(f'学期名称 "{name}" 已存在')
        
        start_date = data.get('start_date')
        end_date = data.get('end_date')
        
        if start_date and end_date and start_date > end_date:
            raise ValidationError('开始日期不能晚于结束日期')
        
        is_current = data.get('is_current', True)
        
        if is_current:
            existing_current = Semester.objects.filter(is_current=True, is_archived=False).first()
            if existing_current:
                raise ValidationError(f'当前已有正在进行的学期「{existing_current.name}」，请先将其归档后再设置新的当前学期')
        
        semester = Semester.objects.create(
            name=name,
            code=code,
            start_date=start_date,
            end_date=end_date,
            total_weeks=data.get('total_weeks', 18),
            description=data.get('description', ''),
            is_current=is_current,
        )

        CacheInvalidator.invalidate_semester_cache()

        return semester

    @transaction.atomic
    def update_semester(self, requester, semester_id: int, data: dict) -> Semester:
        if not requester.is_super_admin:
            raise PermissionDenied('无权限修改学期')
        
        try:
            semester = Semester.objects.get(id=semester_id, is_deleted=False)
        except Semester.DoesNotExist:
            raise NotFoundError('学期不存在')
        
        if semester.is_archived:
            raise ValidationError('已归档的学期无法修改')
        
        if 'name' in data:
            name = data['name'].strip()
            if not name:
                raise ValidationError('学期名称不能为空')
            if Semester.objects.filter(name=name).exclude(id=semester_id).exists():
                raise ValidationError(f'学期名称 "{name}" 已存在')
            semester.name = name
        
        if 'code' in data:
            code = data['code'].strip()
            if not code:
                raise ValidationError('学期代码不能为空')
            if Semester.objects.filter(code=code).exclude(id=semester_id).exists():
                raise ValidationError(f'学期代码 "{code}" 已存在')
            semester.code = code
        
        for field in ['start_date', 'end_date', 'total_weeks', 'description']:
            if field in data:
                setattr(semester, field, data[field])
        
        if semester.start_date and semester.end_date and semester.start_date > semester.end_date:
            raise ValidationError('开始日期不能晚于结束日期')

        semester.save()

        CacheInvalidator.invalidate_semester_cache()

        return semester

    @transaction.atomic
    def delete_semester(self, requester, semester_id: int) -> bool:
        if not requester.is_super_admin:
            raise PermissionDenied('无权限删除学期')
        
        try:
            semester = Semester.objects.get(id=semester_id, is_deleted=False)
        except Semester.DoesNotExist:
            raise NotFoundError('学期不存在')
        
        if semester.is_current:
            raise ValidationError('不能删除当前学期')
        
        if semester.is_archived:
            raise ValidationError('已归档的学期禁止删除')

        semester.delete()

        CacheInvalidator.invalidate_semester_cache()

        return True

    @transaction.atomic
    def batch_delete_semesters(self, requester, semester_ids: list) -> dict:
        if not requester.is_super_admin:
            raise PermissionDenied('无权限删除学期')
        
        deleted_count = 0
        failed_list = []
        
        for semester_id in semester_ids:
            try:
                self.delete_semester(requester, semester_id)
                deleted_count += 1
            except Exception as e:
                failed_list.append({'id': semester_id, 'reason': str(e)})
        
        return {
            'deleted_count': deleted_count,
            'failed_count': len(failed_list),
            'failed_list': failed_list,
        }

    @transaction.atomic
    def set_current_semester(self, requester, semester_id: int) -> Semester:
        if not requester.is_super_admin:
            raise PermissionDenied('无权限设置当前学期')
        
        try:
            semester = Semester.objects.get(id=semester_id, is_deleted=False)
        except Semester.DoesNotExist:
            raise NotFoundError('学期不存在')
        
        if semester.is_archived:
            raise ValidationError('不能将已归档的学期设为当前学期')

        semester.set_as_current()

        CacheInvalidator.invalidate_semester_cache()

        return semester

    @transaction.atomic
    def unset_current_semester(self, requester, semester_id: int) -> Semester:
        if not requester.is_super_admin:
            raise PermissionDenied('无权限操作')
        
        try:
            semester = Semester.objects.get(id=semester_id, is_deleted=False)
        except Semester.DoesNotExist:
            raise NotFoundError('学期不存在')
        
        if not semester.is_current:
            raise ValidationError('该学期不是当前学期')

        semester.is_current = False
        semester.save(update_fields=['is_current'])

        CacheInvalidator.invalidate_semester_cache()

        return semester

    @transaction.atomic
    def archive_semester(self, requester, semester_id: int) -> Semester:
        if not requester.is_super_admin:
            raise PermissionDenied('无权限归档学期')
        
        try:
            semester = Semester.objects.get(id=semester_id, is_deleted=False)
        except Semester.DoesNotExist:
            raise NotFoundError('学期不存在')
        
        if semester.is_current:
            raise ValidationError('不能归档当前学期')
        
        if semester.is_archived:
            raise ValidationError('该学期已归档')

        semester.archive()

        CacheInvalidator.invalidate_semester_cache()

        return semester

    def get_archive_overview(self, requester, department_id: int = None) -> dict:
        if not requester.is_super_admin:
            raise PermissionDenied('无权限查看归档概览')
        
        current_semester = Semester.get_current()
        current_semester_data = None
        
        if current_semester:
            current_semester_data = {
                'id': current_semester.id,
                'name': current_semester.name,
                'start_date': current_semester.start_date.strftime('%Y-%m-%d'),
                'end_date': current_semester.end_date.strftime('%Y-%m-%d'),
                'stats': self._get_semester_stats(current_semester, department_id),
            }
        
        archived_semesters = Semester.objects.filter(is_archived=True).order_by('-end_date')
        archived_list = []
        
        for sem in archived_semesters:
            archived_list.append({
                'id': sem.id,
                'name': sem.name,
                'start_date': sem.start_date.strftime('%Y-%m-%d'),
                'end_date': sem.end_date.strftime('%Y-%m-%d'),
                'stats': self._get_semester_stats(sem, department_id),
            })
        
        departments = self._get_departments()
        
        return {
            'current_semester': current_semester_data,
            'archived_semesters': archived_list,
            'departments': departments,
        }
    
    def _get_departments(self) -> list:
        from apps.users.models import Department
        departments = Department.objects.filter(is_deleted=False).order_by('name')
        return [{'id': d.id, 'name': d.name} for d in departments]

    def _get_semester_stats(self, semester: Semester, department_id: int = None) -> dict:
        lab_filter = {'is_deleted': False}
        if department_id:
            lab_filter['department_id'] = department_id
        
        lab_ids = Laboratory.objects.filter(**lab_filter).values_list('id', flat=True)
        
        return {
            'record_count': UsageRecord.objects.filter(
                semester=semester, laboratory_id__in=lab_ids
            ).count() if lab_ids else 0,
            'maintain_count': WorkOrder.objects.filter(
                semester=semester, laboratory_id__in=lab_ids
            ).exclude(maintenance_type=3).count() if lab_ids else 0,
            'equipment_maintenance_count': WorkOrder.objects.filter(
                semester=semester, laboratory_id__in=lab_ids, maintenance_type=3
            ).count() if lab_ids else 0,
            'class_count': Schedule.objects.filter(
                semester=semester, laboratory_id__in=lab_ids
            ).count() if lab_ids else 0,
            'lab_count': len(lab_ids),
            'device_count': Equipment.objects.filter(
                laboratory_id__in=lab_ids, is_deleted=False
            ).count() if lab_ids else 0,
            'user_count': User.objects.filter(
                department_id=department_id, is_deleted=False
            ).count() if department_id else User.objects.filter(is_deleted=False).count(),
        }

    def _format_semester(self, semester: Semester) -> dict:
        return {
            'id': semester.id,
            'name': semester.name,
            'code': semester.code,
            'start_date': semester.start_date.strftime('%Y-%m-%d'),
            'end_date': semester.end_date.strftime('%Y-%m-%d'),
            'is_current': semester.is_current,
            'is_archived': semester.is_archived,
            'total_weeks': semester.total_weeks,
            'schedule_count': semester.get_schedule_count(),
            'created_at': semester.created_at.strftime('%Y-%m-%d %H:%M:%S'),
        }

    def _format_semester_detail(self, semester: Semester) -> dict:
        data = self._format_semester(semester)
        data.update({
            'description': semester.description,
            'updated_at': semester.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
        })
        return data

    def get_archive_settings(self, requester) -> dict:
        if not requester.is_super_admin:
            raise PermissionDenied('无权限访问')

        cache_key = CacheKeyManager.make_key(CacheKeyManager.ARCHIVE_SETTINGS)
        cached = cache.get(cache_key)
        if cached is not None:
            return {'retention_months': cached}

        setting, created = SystemSetting.objects.get_or_create(
            key='archive_retention_months',
            defaults={'value': 0, 'description': '归档保留时间(月)，0表示永久保留'}
        )
        cache.set(cache_key, setting.value, 3600)
        return {'retention_months': setting.value}

    @transaction.atomic
    def save_archive_settings(self, requester, retention_months: int) -> dict:
        if not requester.is_super_admin:
            raise PermissionDenied('无权限访问')

        try:
            retention_months = int(retention_months)
            if retention_months < 0:
                raise ValueError
        except (TypeError, ValueError):
            raise ValidationError('请输入有效的月份数值')

        SystemSetting.objects.update_or_create(
            key='archive_retention_months',
            defaults={
                'value': retention_months,
                'description': '归档保留时间(月)，0表示永久保留'
            }
        )

        cache_key = CacheKeyManager.make_key(CacheKeyManager.ARCHIVE_SETTINGS)
        cache.set(cache_key, retention_months, 3600)

        cleanup_msg = ''
        if retention_months > 0:
            deleted_count = self._cleanup_expired_archives(retention_months)
            if deleted_count > 0:
                cleanup_msg = f'，并自动清理了 {deleted_count} 个已过期的归档学期'

        return {'retention_months': retention_months, 'message': f'设置已保存{cleanup_msg}'}

    def manual_cleanup(self, requester) -> dict:
        if not requester.is_super_admin:
            raise PermissionDenied('只有超级管理员可以执行清理操作')

        cache_key = CacheKeyManager.make_key(CacheKeyManager.ARCHIVE_SETTINGS)
        cached_retention = cache.get(cache_key)
        if cached_retention is not None:
            retention_months = cached_retention
        else:
            try:
                setting = SystemSetting.objects.get(key='archive_retention_months')
                retention_months = int(setting.value)
                cache.set(cache_key, retention_months, 3600)
            except SystemSetting.DoesNotExist:
                raise ValidationError('请先设置过期时间')

        if retention_months <= 0:
            raise ValidationError('请先设置过期时间')

        deleted_count = self._cleanup_expired_archives(retention_months)

        if deleted_count > 0:
            return {'deleted_count': deleted_count, 'message': f'成功清理了 {deleted_count} 个过期的归档学期'}
        else:
            return {'deleted_count': 0, 'message': '没有找到需要清理的过期归档学期'}

    def _cleanup_expired_archives(self, months: int) -> int:
        if months <= 0:
            return 0

        cutoff_date = timezone.now().date() - timedelta(days=months * 30)

        with transaction.atomic():
            expired_ids = list(Semester.objects.filter(
                is_archived=True,
                end_date__lt=cutoff_date
            ).values_list('id', flat=True))

            count = len(expired_ids)
            if count > 0:
                Semester.objects.filter(id__in=expired_ids).delete()
                cache_key = CacheKeyManager.make_key(CacheKeyManager.ARCHIVE_SETTINGS)
                cache.delete(cache_key)

        return count

    @transaction.atomic
    def archive_current_term_records(self, requester, archive_types: list, department_id: int = None) -> dict:
        if not requester.is_super_admin:
            raise PermissionDenied('只有超级管理员可以归档学期记录')
        
        if not archive_types:
            raise ValidationError('请至少选择一种要归档的数据类型')
        
        current_semester = Semester.objects.select_for_update().filter(
            is_current=True, is_archived=False
        ).first()
        
        if not current_semester:
            raise ValidationError('未设置当前学期，请先由超级管理员创建并设为当前学期')
        
        if current_semester.is_archived:
            raise ValidationError('当前学期记录已经归档')
        
        archived_items = []
        archived_counts = {}
        
        TermArchive.objects.filter(
            semester=current_semester,
            archive_type__in=archive_types
        ).delete()
        
        if 'lab_info' in archive_types:
            TermArchive.objects.filter(
                semester=current_semester,
                archive_type__in=['lab_info', 'device_info', 'usage_records', 'schedules', 'maintain_records', 'fault_records']
            ).delete()
        
        lab_filter = {'is_deleted': False}
        if department_id:
            lab_filter['department_id'] = department_id
        lab_ids = list(Laboratory.objects.filter(**lab_filter).values_list('id', flat=True))
        
        if 'lab_info' in archive_types:
            labs = Laboratory.objects.filter(**lab_filter)
            if labs.exists():
                lab_count = labs.count()
                equipment_count = 0
                record_count = 0
                schedule_count = 0
                maintain_count = 0
                fault_count = 0
                
                lab_data = []
                for lab in labs:
                    lab_dict = {
                        'pk': lab.id,
                        'fields': {
                            'name': lab.name,
                            'code': lab.code,
                            'building': lab.building,
                            'floor': lab.floor,
                            'room_number': lab.room_number,
                            'capacity': lab.capacity,
                            'area': lab.area,
                            'laboratory_type': lab.laboratory_type,
                            'department_id': lab.department_id,
                            'admin_id': lab.admin_id,
                            'status': lab.status,
                            'is_available': lab.is_available,
                            'facilities': lab.facilities,
                            'description': lab.description,
                            'note': lab.note,
                            'images': lab.images,
                        }
                    }
                    lab_data.append(lab_dict)
                
                TermArchive.objects.create(
                    semester=current_semester,
                    archive_type='lab_info',
                    content=lab_data
                )
                
                equipments = Equipment.objects.filter(laboratory_id__in=lab_ids, is_deleted=False) if lab_ids else Equipment.objects.none()
                if equipments.exists():
                    equipment_data = []
                    for eq in equipments:
                        eq_dict = {
                            'pk': eq.id,
                            'fields': {
                                'name': eq.name,
                                'code': eq.code,
                                'category': eq.category,
                                'brand': eq.brand,
                                'model': eq.model,
                                'laboratory_id': eq.laboratory_id,
                                'status': eq.status,
                                'purchase_date': str(eq.purchase_date) if eq.purchase_date else None,
                                'price': str(eq.price) if eq.price else None,
                                'description': eq.description,
                            }
                        }
                        equipment_data.append(eq_dict)
                    equipment_count = len(equipment_data)
                    TermArchive.objects.create(
                        semester=current_semester,
                        archive_type='device_info',
                        content=equipment_data
                    )
                
                records = UsageRecord.objects.filter(semester=current_semester, laboratory_id__in=lab_ids) if lab_ids else UsageRecord.objects.none()
                if records.exists():
                    record_data = []
                    for r in records:
                        r_dict = {
                            'pk': r.id,
                            'fields': {
                                'usage_date': str(r.usage_date),
                                'time_slot': r.time_slot,
                                'class_hours': r.class_hours,
                                'laboratory_id': r.laboratory_id,
                                'laboratory_name': r.laboratory_name,
                                'laboratory_code': r.laboratory_code,
                                'semester_id': r.semester_id,
                                'teacher_id': r.teacher_id,
                                'class_name': r.class_name,
                                'student_count': r.student_count,
                                'content': r.content,
                                'device_status': r.device_status,
                                'laboratory_status': r.laboratory_status,
                                'note': r.note,
                            }
                        }
                        record_data.append(r_dict)
                    record_count = len(record_data)
                    TermArchive.objects.create(
                        semester=current_semester,
                        archive_type='usage_records',
                        content=record_data
                    )
                
                schedules = Schedule.objects.filter(semester=current_semester, laboratory_id__in=lab_ids) if lab_ids else Schedule.objects.none()
                if schedules.exists():
                    schedule_data = []
                    for s in schedules:
                        s_dict = {
                            'pk': s.id,
                            'fields': {
                                'course_name': s.course_name,
                                'course_code': s.course_code,
                                'weekday': s.weekday,
                                'time_slot': s.time_slot,
                                'weeks': s.weeks,
                                'laboratory_id': s.laboratory_id,
                                'semester_id': s.semester_id,
                                'teacher_id': s.teacher_id,
                                'teacher_name': s.teacher_name,
                                'class_name': s.class_name,
                                'student_count': s.student_count,
                                'note': s.note,
                            }
                        }
                        schedule_data.append(s_dict)
                    schedule_count = len(schedule_data)
                    TermArchive.objects.create(
                        semester=current_semester,
                        archive_type='schedules',
                        content=schedule_data
                    )
                
                workorders = WorkOrder.objects.filter(semester=current_semester, laboratory_id__in=lab_ids) if lab_ids else WorkOrder.objects.none()
                if workorders.exists():
                    maintain_data = []
                    fault_data = []
                    for w in workorders:
                        w_dict = {
                            'pk': w.id,
                            'fields': {
                                'order_number': w.order_number,
                                'title': w.title,
                                'description': w.description,
                                'laboratory_id': w.laboratory_id,
                                'laboratory_name': w.laboratory_name,
                                'laboratory_code': w.laboratory_code,
                                'equipment_id': w.equipment_id,
                                'semester_id': w.semester_id,
                                'maintenance_type': w.maintenance_type,
                                'status': w.status,
                                'priority': w.priority,
                                'reporter_id': w.reporter_id,
                                'handler_id': w.handler_id,
                                'reported_at': str(w.reported_at) if w.reported_at else None,
                                'assigned_at': str(w.assigned_at) if w.assigned_at else None,
                                'started_at': str(w.started_at) if w.started_at else None,
                                'completed_at': str(w.completed_at) if w.completed_at else None,
                                'closed_at': str(w.closed_at) if w.closed_at else None,
                                'solution': w.solution,
                                'handle_note': w.handle_note,
                                'rating': w.rating,
                                'feedback': w.feedback,
                            }
                        }
                        if w.maintenance_type == 3:
                            fault_data.append(w_dict)
                        else:
                            maintain_data.append(w_dict)
                    
                    maintain_count = len(maintain_data)
                    fault_count = len(fault_data)
                    
                    if maintain_data:
                        TermArchive.objects.create(
                            semester=current_semester,
                            archive_type='maintain_records',
                            content=maintain_data
                        )
                    if fault_data:
                        TermArchive.objects.create(
                            semester=current_semester,
                            archive_type='fault_records',
                            content=fault_data
                        )
                
                records.delete()
                schedules.delete()
                workorders.delete()
                
                labs.delete()
                
                archived_items.append('实训室信息')
                archived_counts['实训室信息'] = lab_count
                if equipment_count > 0:
                    archived_items.append('设备信息')
                    archived_counts['设备信息'] = equipment_count
                if record_count > 0:
                    archived_items.append('使用记录')
                    archived_counts['使用记录'] = record_count
                if schedule_count > 0:
                    archived_items.append('课表记录')
                    archived_counts['课表记录'] = schedule_count
                if maintain_count > 0:
                    archived_items.append('维护记录')
                    archived_counts['维护记录'] = maintain_count
                if fault_count > 0:
                    archived_items.append('故障记录')
                    archived_counts['故障记录'] = fault_count
        
        elif 'device_info' in archive_types:
            equipment = Equipment.objects.filter(laboratory_id__in=lab_ids, is_deleted=False) if lab_ids else Equipment.objects.filter(is_deleted=False)
            if equipment.exists():
                equipment_data = []
                for eq in equipment:
                    eq_dict = {
                        'pk': eq.id,
                        'fields': {
                            'name': eq.name,
                            'code': eq.code,
                            'category': eq.category,
                            'brand': eq.brand,
                            'model': eq.model,
                            'laboratory_id': eq.laboratory_id,
                            'status': eq.status,
                            'purchase_date': str(eq.purchase_date) if eq.purchase_date else None,
                            'price': str(eq.price) if eq.price else None,
                            'description': eq.description,
                        }
                    }
                    equipment_data.append(eq_dict)
                TermArchive.objects.create(
                    semester=current_semester,
                    archive_type='device_info',
                    content=equipment_data
                )
                count = len(equipment_data)
                equipment.delete()
                archived_items.append('设备信息')
                archived_counts['设备信息'] = count
        
        if 'user_info' in archive_types:
            user_filter = {'is_deleted': False}
            if department_id:
                user_filter['department_id'] = department_id
            users = User.objects.filter(**user_filter)
            if users.exists():
                user_data = []
                for u in users:
                    u_dict = {
                        'pk': u.id,
                        'fields': {
                            'username': u.username,
                            'nickname': u.nickname,
                            'department_id': u.department_id,
                            'role': u.role,
                            'phone': u.phone,
                            'email': u.email,
                        }
                    }
                    user_data.append(u_dict)
                TermArchive.objects.create(
                    semester=current_semester,
                    archive_type='user_info',
                    content=user_data
                )
                archived_items.append('用户信息')
                archived_counts['用户信息'] = len(user_data)
        
        if archived_items:
            if not department_id:
                current_semester.is_archived = True
                current_semester.is_current = False
                current_semester.save(update_fields=['is_archived', 'is_current'])
        
        message_parts = []
        for item in archived_items:
            count = archived_counts.get(item, 0)
            message_parts.append(f'{item}({count}条)')
        
        dept_msg = f'部门ID:{department_id}的' if department_id else ''
        message = f'成功归档学期「{current_semester.name}」{dept_msg}以下数据：{", ".join(message_parts)}。归档后除超级管理员外其他用户只能查看不能修改'
        
        return {
            'archived_items': archived_items,
            'archived_counts': archived_counts,
            'semester_id': current_semester.id,
            'message': message,
        }

    def get_archived_records(self, requester, semester_id: int, record_type: str = None,
                             department_id: int = None, page: int = 1, page_size: int = 10,
                             export_format: str = None) -> dict:
        if not requester.is_super_admin and not requester.is_department_admin:
            raise PermissionDenied('只有超级管理员或分院管理员可以查看归档记录')
        
        if requester.is_department_admin and not requester.is_super_admin:
            department_id = requester.department_id
        
        try:
            semester = Semester.objects.get(id=semester_id)
        except Semester.DoesNotExist:
            raise NotFoundError('学期不存在')
        
        if export_format in ['excel', 'pdf']:
            return self._export_archived_records(semester, export_format, department_id)
        
        if not record_type:
            stats = self._get_archived_dashboard_stats(semester, department_id)
            return {
                'mode': 'dashboard',
                'semester_name': semester.name,
                'stats': stats,
            }
        
        records_list = self._get_archived_records_by_type(semester, record_type, department_id)
        
        paginator = Paginator(records_list, page_size)
        page_obj = paginator.get_page(page)
        
        return {
            'mode': 'list',
            'semester_name': semester.name,
            'record_type': record_type,
            'list': list(page_obj),
            'total': paginator.count,
            'page': page,
            'page_size': page_size,
            'total_pages': paginator.num_pages,
        }

    def _get_archived_dashboard_stats(self, semester: Semester, department_id: int = None) -> dict:
        lab_archive = TermArchive.objects.filter(semester=semester, archive_type='lab_info').first()
        lab_count = len(lab_archive.content) if lab_archive else 0
        
        device_archive = TermArchive.objects.filter(semester=semester, archive_type='device_info').first()
        device_count = len(device_archive.content) if device_archive else 0
        
        user_archive = TermArchive.objects.filter(semester=semester, archive_type='user_info').first()
        user_count = len(user_archive.content) if user_archive else 0
        
        usage_archive = TermArchive.objects.filter(semester=semester, archive_type='usage_records').first()
        usage_count = len(usage_archive.content) if usage_archive else 0
        
        schedule_archive = TermArchive.objects.filter(semester=semester, archive_type='schedules').first()
        schedule_count = len(schedule_archive.content) if schedule_archive else 0
        
        maintain_archive = TermArchive.objects.filter(semester=semester, archive_type='maintain_records').first()
        maintain_count = len(maintain_archive.content) if maintain_archive else 0
        
        fault_archive = TermArchive.objects.filter(semester=semester, archive_type='fault_records').first()
        fault_count = len(fault_archive.content) if fault_archive else 0
        
        return {
            'record_count': usage_count,
            'maintain_count': maintain_count,
            'equipment_maintenance_count': fault_count,
            'class_count': schedule_count,
            'lab_count': lab_count,
            'device_count': device_count,
            'user_count': user_count,
            'usage': usage_count,
            'maintain': maintain_count,
            'fault': fault_count,
            'class': schedule_count,
            'lab': lab_count,
            'device': device_count,
            'user': user_count,
        }

    def _get_archived_records_by_type(self, semester: Semester, record_type: str,
                                       department_id: int = None) -> list:
        records_list = []
        
        lab_archive = TermArchive.objects.filter(semester=semester, archive_type='lab_info').first()
        lab_map = {}
        if lab_archive:
            for item in lab_archive.content:
                lab_map[item['pk']] = item['fields'].get('name', '暂无')
        
        user_archive = TermArchive.objects.filter(semester=semester, archive_type='user_info').first()
        user_map = {}
        if user_archive:
            for item in user_archive.content:
                user_map[item['pk']] = item['fields'].get('nickname', '') or item['fields'].get('username', '暂无')
        
        from apps.users.models import Department
        dept_map = {}
        depts = Department.objects.filter(is_deleted=False)
        for dept in depts:
            dept_map[dept.id] = dept.name
        
        status_map = {
            1: '待处理', 2: '处理中', 3: '已完成', 4: '已关闭', 5: '已取消'
        }
        work_order_status_map = {
            'PENDING': '待处理', 'ASSIGNED': '已分配',
            'IN_PROGRESS': '处理中', 'COMPLETED': '已完成',
            'CLOSED': '已关闭', 'CANCELLED': '已取消'
        }
        
        if record_type == 'usage':
            archive = TermArchive.objects.filter(semester=semester, archive_type='usage_records').first()
            if archive:
                for item in archive.content:
                    fields = item['fields']
                    lab_id = fields.get('laboratory_id')
                    teacher_id = fields.get('teacher_id')
                    content_val = fields.get('content', '') or ''
                    records_list.append({
                        'id': item['pk'],
                        'laboratory_name': lab_map.get(lab_id, '暂无'),
                        'class_name': fields.get('class_name') or '暂无',
                        'teacher_name': user_map.get(teacher_id, '暂无'),
                        'content': content_val[:30] + '...' if len(content_val) > 30 else content_val,
                        'usage_date': fields.get('usage_date') or '暂无',
                        'time_slot': fields.get('time_slot') or '暂无',
                        'operator': user_map.get(teacher_id, '暂无'),
                    })
        
        elif record_type == 'maintain':
            archive = TermArchive.objects.filter(semester=semester, archive_type='maintain_records').first()
            if archive:
                for item in archive.content:
                    fields = item['fields']
                    lab_id = fields.get('laboratory_id')
                    reporter_id = fields.get('reporter_id')
                    handler_id = fields.get('handler_id')
                    description = fields.get('description', '') or ''
                    reported_at = fields.get('reported_at', '') or ''
                    completed_at = fields.get('completed_at', '') or ''
                    status_val = fields.get('status')
                    status_text = work_order_status_map.get(status_val, str(status_val) if status_val else '暂无')
                    records_list.append({
                        'id': item['pk'],
                        'order_number': fields.get('order_number') or f'WO-{item["pk"]}',
                        'title': fields.get('title') or f'维护记录-{item["pk"]}',
                        'laboratory_name': lab_map.get(lab_id, '暂无'),
                        'reporter_name': user_map.get(reporter_id, '暂无'),
                        'handler_name': user_map.get(handler_id, '暂无'),
                        'description': description[:30] + '...' if len(description) > 30 else description,
                        'status_display': status_text,
                        'status': status_val,
                        'reported_at': reported_at[:10] if reported_at else '暂无',
                        'completed_at': completed_at[:10] if completed_at else '暂无',
                        'priority': fields.get('priority') or '暂无',
                    })
        
        elif record_type == 'fault':
            archive = TermArchive.objects.filter(semester=semester, archive_type='fault_records').first()
            if archive:
                for item in archive.content:
                    fields = item['fields']
                    lab_id = fields.get('laboratory_id')
                    reporter_id = fields.get('reporter_id')
                    handler_id = fields.get('handler_id')
                    description = fields.get('description', '') or ''
                    reported_at = fields.get('reported_at', '') or ''
                    completed_at = fields.get('completed_at', '') or ''
                    status_val = fields.get('status')
                    status_text = work_order_status_map.get(status_val, str(status_val) if status_val else '暂无')
                    records_list.append({
                        'id': item['pk'],
                        'order_number': fields.get('order_number') or f'WO-{item["pk"]}',
                        'title': fields.get('title') or f'故障工单-{item["pk"]}',
                        'laboratory_name': lab_map.get(lab_id, '暂无'),
                        'reporter_name': user_map.get(reporter_id, '暂无'),
                        'handler_name': user_map.get(handler_id, '暂无'),
                        'description': description[:30] + '...' if len(description) > 30 else description,
                        'status_display': status_text,
                        'status': status_val,
                        'reported_at': reported_at[:10] if reported_at else '暂无',
                        'completed_at': completed_at[:10] if completed_at else '暂无',
                        'priority': fields.get('priority') or '暂无',
                    })
        
        elif record_type == 'class':
            archive = TermArchive.objects.filter(semester=semester, archive_type='schedules').first()
            if archive:
                for item in archive.content:
                    fields = item['fields']
                    lab_id = fields.get('laboratory_id')
                    teacher_id = fields.get('teacher_id')
                    records_list.append({
                        'id': item['pk'],
                        'course_name': fields.get('course_name') or '暂无',
                        'laboratory_name': lab_map.get(lab_id, '暂无'),
                        'teacher_name': user_map.get(teacher_id, fields.get('teacher_name')) or '暂无',
                        'class_name': fields.get('class_name') or '暂无',
                        'weekday': fields.get('weekday') or '暂无',
                        'time_slot': fields.get('time_slot') or '暂无',
                        'weeks': fields.get('weeks') or '暂无',
                    })
        
        elif record_type == 'lab_info':
            archive = TermArchive.objects.filter(semester=semester, archive_type='lab_info').first()
            if archive:
                for item in archive.content:
                    fields = item['fields']
                    admin_id = fields.get('admin_id')
                    dept_id = fields.get('department_id')
                    room_number = fields.get('room_number', '') or ''
                    building = fields.get('building', '') or ''
                    floor = str(fields.get('floor', '')) if fields.get('floor') else ''
                    room_display = f"{building}{floor}层{room_number}" if building or floor or room_number else '暂无'
                    records_list.append({
                        'id': item['pk'],
                        'name': fields.get('name') or '暂无',
                        'code': fields.get('code') or '暂无',
                        'room_number': room_display,
                        'admin_name': user_map.get(admin_id, '暂无'),
                        'department_name': dept_map.get(dept_id, '暂无'),
                    })
        
        elif record_type == 'device_info':
            archive = TermArchive.objects.filter(semester=semester, archive_type='device_info').first()
            if archive:
                for item in archive.content:
                    fields = item['fields']
                    loc_id = fields.get('laboratory_id')
                    loc_name = lab_map.get(loc_id, '未分配') if loc_id else '未分配'
                    description = fields.get('description', '') or ''
                    
                    records_list.append({
                        'id': item['pk'],
                        'code': fields.get('code') or '暂无',
                        'name': fields.get('name') or '暂无',
                        'brand': fields.get('brand') or '暂无',
                        'model': fields.get('model') or '暂无',
                        'category': fields.get('category') or '暂无',
                        'description': description[:30] + '...' if len(description) > 30 else description,
                        'laboratory_name': loc_name,
                    })
        
        elif record_type == 'user_info':
            archive = TermArchive.objects.filter(semester=semester, archive_type='user_info').first()
            if archive:
                role_map = {
                    1: '教师', 2: '实训室管理员', 4: '分院管理员', 16: '超级管理员', 32: '系统管理员'
                }
                for item in archive.content:
                    fields = item['fields']
                    role_val = fields.get('role', 0) or 0
                    dept_id = fields.get('department_id')
                    
                    roles = []
                    for r_val, r_name in role_map.items():
                        if role_val & r_val:
                            roles.append(r_name)
                    role_display = '、'.join(roles) if roles else '暂无'
                    
                    records_list.append({
                        'id': item['pk'],
                        'username': fields.get('username') or '暂无',
                        'email': fields.get('email') or '暂无',
                        'phone': fields.get('phone') or '暂无',
                        'nickname': fields.get('nickname') or '暂无',
                        'role_display': role_display,
                        'department_name': dept_map.get(dept_id, '暂无'),
                    })
        
        return records_list

    def _export_archived_records(self, semester: Semester, format_type: str,
                                  department_id: int = None):
        from openpyxl import Workbook

        if format_type == 'excel':
            wb = Workbook()
            default_ws = wb.active
            wb.remove(default_ws)

            lab_archive = TermArchive.objects.filter(semester=semester, archive_type='lab_info').first()
            lab_map = {}
            if lab_archive:
                for item in lab_archive.content:
                    lab_map[item['pk']] = item['fields'].get('name', '暂无')

            user_archive = TermArchive.objects.filter(semester=semester, archive_type='user_info').first()
            user_map = {}
            if user_archive:
                for item in user_archive.content:
                    user_map[item['pk']] = item['fields'].get('nickname', '') or item['fields'].get('username', '暂无')

            usage_rows = []
            usage_archive = TermArchive.objects.filter(semester=semester, archive_type='usage_records').first()
            if usage_archive:
                for item in usage_archive.content:
                    fields = item['fields']
                    lab_id = fields.get('laboratory_id')
                    teacher_id = fields.get('teacher_id')
                    usage_rows.append([
                        item['pk'],
                        lab_map.get(lab_id, '暂无'),
                        fields.get('content', '') or '',
                        fields.get('usage_date') or '',
                        user_map.get(teacher_id, '暂无'),
                        fields.get('class_name') or '',
                    ])

            if usage_rows:
                ws = wb.create_sheet(title='使用记录')
                headers = ['ID', '实训室', '内容', '日期', '教师', '班级']
                for col, header in enumerate(headers, 1):
                    ws.cell(row=1, column=col, value=header)
                for row_idx, row_data in enumerate(usage_rows, 2):
                    for col_idx, value in enumerate(row_data, 1):
                        ws.cell(row=row_idx, column=col_idx, value=str(value) if value else '')

            schedule_rows = []
            schedule_archive = TermArchive.objects.filter(semester=semester, archive_type='schedules').first()
            if schedule_archive:
                for item in schedule_archive.content:
                    fields = item['fields']
                    lab_id = fields.get('laboratory_id')
                    teacher_id = fields.get('teacher_id')
                    schedule_rows.append([
                        item['pk'],
                        fields.get('course_name') or '',
                        lab_map.get(lab_id, '暂无'),
                        user_map.get(teacher_id, fields.get('teacher_name')) or '',
                        fields.get('class_name') or '',
                        fields.get('weekday') or '',
                        fields.get('time_slot') or '',
                        fields.get('weeks') or '',
                    ])

            if schedule_rows:
                ws = wb.create_sheet(title='课表记录')
                headers = ['ID', '课程名称', '实训室', '教师', '班级', '星期', '节次', '周次']
                for col, header in enumerate(headers, 1):
                    ws.cell(row=1, column=col, value=header)
                for row_idx, row_data in enumerate(schedule_rows, 2):
                    for col_idx, value in enumerate(row_data, 1):
                        ws.cell(row=row_idx, column=col_idx, value=str(value) if value else '')

            maintain_rows = []
            maintain_archive = TermArchive.objects.filter(semester=semester, archive_type='maintain_records').first()
            if maintain_archive:
                status_map = {
                    'PENDING': '待处理', 'ASSIGNED': '已分配',
                    'IN_PROGRESS': '处理中', 'COMPLETED': '已完成',
                    'CLOSED': '已关闭', 'CANCELLED': '已取消'
                }
                for item in maintain_archive.content:
                    fields = item['fields']
                    lab_id = fields.get('laboratory_id')
                    reporter_id = fields.get('reporter_id')
                    handler_id = fields.get('handler_id')
                    status_val = fields.get('status')
                    maintain_rows.append([
                        item['pk'],
                        fields.get('title') or f'维护记录-{item["pk"]}',
                        lab_map.get(lab_id, '暂无'),
                        user_map.get(reporter_id, '暂无'),
                        user_map.get(handler_id, '暂无'),
                        status_map.get(status_val, str(status_val) if status_val else '暂无'),
                        fields.get('reported_at', '')[:10] if fields.get('reported_at') else '',
                        fields.get('completed_at', '')[:10] if fields.get('completed_at') else '',
                    ])

            if maintain_rows:
                ws = wb.create_sheet(title='维护记录')
                headers = ['ID', '标题', '实训室', '报告人', '处理人', '状态', '报告时间', '完成时间']
                for col, header in enumerate(headers, 1):
                    ws.cell(row=1, column=col, value=header)
                for row_idx, row_data in enumerate(maintain_rows, 2):
                    for col_idx, value in enumerate(row_data, 1):
                        ws.cell(row=row_idx, column=col_idx, value=str(value) if value else '')

            fault_rows = []
            fault_archive = TermArchive.objects.filter(semester=semester, archive_type='fault_records').first()
            if fault_archive:
                status_map = {
                    'PENDING': '待处理', 'ASSIGNED': '已分配',
                    'IN_PROGRESS': '处理中', 'COMPLETED': '已完成',
                    'CLOSED': '已关闭', 'CANCELLED': '已取消'
                }
                for item in fault_archive.content:
                    fields = item['fields']
                    lab_id = fields.get('laboratory_id')
                    reporter_id = fields.get('reporter_id')
                    handler_id = fields.get('handler_id')
                    status_val = fields.get('status')
                    fault_rows.append([
                        item['pk'],
                        fields.get('title') or f'故障工单-{item["pk"]}',
                        lab_map.get(lab_id, '暂无'),
                        user_map.get(reporter_id, '暂无'),
                        user_map.get(handler_id, '暂无'),
                        status_map.get(status_val, str(status_val) if status_val else '暂无'),
                        fields.get('reported_at', '')[:10] if fields.get('reported_at') else '',
                        fields.get('completed_at', '')[:10] if fields.get('completed_at') else '',
                    ])

            if fault_rows:
                ws = wb.create_sheet(title='故障工单')
                headers = ['ID', '标题', '实训室', '报告人', '处理人', '状态', '报告时间', '完成时间']
                for col, header in enumerate(headers, 1):
                    ws.cell(row=1, column=col, value=header)
                for row_idx, row_data in enumerate(fault_rows, 2):
                    for col_idx, value in enumerate(row_data, 1):
                        ws.cell(row=row_idx, column=col_idx, value=str(value) if value else '')

            if not wb.worksheets:
                raise ValidationError('该学期没有可导出的归档数据')

            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            filename = f"学期{semester.name}归档记录.xlsx"
            response['Content-Disposition'] = f'attachment; filename="{escape_uri_path(filename)}"'
            wb.save(response)
            return response

        raise ValidationError('不支持的导出格式')

    def check_term_status(self, requester) -> dict:
        current_semester = Semester.get_current()
        
        return {
            'is_super_admin': requester.is_super_admin,
            'has_current_semester': current_semester is not None,
            'current_semester': {
                'id': current_semester.id,
                'name': current_semester.name,
            } if current_semester else None,
        }
