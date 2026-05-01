"""
用户服务
"""

import io
from django.db import models, transaction
from django.core.paginator import Paginator
from openpyxl import load_workbook
from apps.core.exceptions import ValidationError, NotFoundError, PermissionDenied
from apps.core.constants import UserRole
from apps.users.models import User, Department
from apps.core.services.operation_log_service import OperationLogService
from common.decorators import cached_method
from common.services.cache_service import CacheInvalidator
from apps.core.utils import beijing_strftime, beijing_now, beijing_today


class UserService:
    """用户服务"""

    @cached_method(timeout=60, key_prefix='user:list')
    def get_user_list(
        self,
        requester,
        department_id: int = None,
        role: int = None,
        status: str = None,
        search: str = None,
        page: int = 1,
        page_size: int = 20,
        no_page: bool = False
    ) -> dict:
        queryset = User.objects.select_related('department').filter(is_deleted=False)
        
        if requester.is_superuser:
            pass
        
        elif requester.is_super_admin:
            pass
        
        elif requester.is_department_admin:
            if not requester.department_id:
                return {
                    'list': [],
                    'pagination': {
                        'total': 0,
                        'page': page,
                        'page_size': page_size,
                        'total_pages': 0,
                    },
                    'warning': 'no_department_assigned',
                    'message': '您还未被分配部门，无法查看用户列表。请联系超级管理员为您分配部门。'
                }
            
            queryset = queryset.filter(department_id=requester.department_id)
            queryset = queryset.exclude(id=requester.id)
            queryset = queryset.exclude(department_id__isnull=True)
        
        if department_id:
            queryset = queryset.filter(department_id=department_id)
        if role is not None:
            from django.db.models import IntegerField, Q
            
            if isinstance(role, str) and ',' in role:
                role_list = [int(r.strip()) for r in role.split(',') if r.strip().isdigit()]
                if role_list:
                    where_conditions = ' OR '.join([f'(role & {r} > 0)' for r in role_list])
                    queryset = queryset.extra(where=[where_conditions])
            else:
                try:
                    role_int = int(role)
                    queryset = queryset.annotate(
                        role_bitand=models.ExpressionWrapper(
                            models.F('role').bitand(role_int),
                            output_field=IntegerField()
                        )
                    ).filter(role_bitand__gt=0)
                except (ValueError, TypeError):
                    pass
        if status:
            queryset = queryset.filter(status=status)
        if search:
            queryset = queryset.filter(
                models.Q(username__icontains=search) |
                models.Q(nickname__icontains=search) |
                models.Q(phone__icontains=search)
            )
        
        queryset = queryset.order_by('-created_at')
        
        if no_page:
            return {
                'list': [self._format_user(u) for u in queryset],
                'pagination': {
                    'total': queryset.count(),
                }
            }
        
        paginator = Paginator(queryset, page_size)
        page_obj = paginator.get_page(page)
        
        return {
            'list': [self._format_user(u) for u in page_obj],
            'pagination': {
                'total': paginator.count,
                'page': page,
                'page_size': page_size,
                'total_pages': paginator.num_pages,
            }
        }

    @transaction.atomic
    def create_user(self, requester, data: dict, request=None):
        if not requester.is_department_admin and not requester.is_super_admin:
            raise PermissionDenied('无权限创建用户')

        username = data.get('username', '').strip()
        if not username:
            raise ValidationError('用户名不能为空')

        if User.objects.filter(username=username).exists():
            raise ValidationError(f'用户名 "{username}" 已存在')

        department = data.get('department')
        
        role = data.get('role', UserRole.TEACHER)
        
        if role & UserRole.DEPARTMENT_ADMIN:
            if not department:
                raise ValidationError('分院管理员必须指定所属部门')
        
        if department:
            if not requester.is_super_admin:
                if department.id != requester.department_id:
                    raise PermissionDenied('无权限在该部门创建用户')
        elif not requester.is_super_admin and not (role & UserRole.TEACHER or role & UserRole.LABORATORY_ADMIN):
            if requester.department_id:
                from apps.users.models import Department
                try:
                    department = Department.objects.get(id=requester.department_id)
                except Department.DoesNotExist:
                    pass

        default_password = username[:6]

        default_role = UserRole.TEACHER
        if requester.is_super_admin and 'role' not in data:
            default_role = UserRole.DEPARTMENT_ADMIN

        user = User.objects.create_user(
            username=username,
            nickname=data.get('nickname', username),
            phone=data.get('phone', ''),
            email=data.get('email', ''),
            role=role,
            department=department,
            is_active=True,
            password=data.get('password', default_password)
        )

        if department and (role & UserRole.DEPARTMENT_ADMIN or role & UserRole.SUPER_ADMIN or role & UserRole.SYSTEM_ADMIN):
            if user.id not in department.managers.values_list('id', flat=True):
                department.managers.add(user)

        CacheInvalidator.invalidate_user_cache(requester.id)

        OperationLogService.log_user_operation(
            request=request,
            operation_type='user_create',
            user=user
        )

        return user

    @transaction.atomic
    def update_user(self, requester, user_id: int, data: dict, request=None):
        try:
            user = User.objects.get(id=user_id, is_deleted=False)
        except User.DoesNotExist:
            raise NotFoundError('用户不存在')
        
        if not self._can_manage_user(requester, user):
            raise PermissionDenied('无权限修改该用户')

        old_data = {
            'nickname': user.nickname,
            'phone': user.phone,
            'email': user.email,
            'status': user.status,
            'role': user.role,
            'department_id': user.department_id,
        }
        
        new_role = data.get('role', user.role)
        new_department_id = data.get('department_id', user.department_id)
        
        if new_role & UserRole.DEPARTMENT_ADMIN:
            if not new_department_id:
                raise ValidationError('分院管理员必须指定所属部门')
        
        allowed_fields = ['nickname', 'phone', 'email', 'status']
        if requester.is_super_admin or requester.is_department_admin:
            allowed_fields.extend(['role', 'department_id'])
        
        for field in allowed_fields:
            if field in data:
                setattr(user, field, data[field])
        
        user.save()

        if 'department_id' in data or 'role' in data:
            from apps.users.models import Department
            
            old_dept_id = old_data['department_id']
            old_role = old_data['role']
            new_dept_id = user.department_id
            new_role = user.role
            
            was_admin = old_role & UserRole.DEPARTMENT_ADMIN or old_role & UserRole.SUPER_ADMIN or old_role & UserRole.SYSTEM_ADMIN
            is_admin = new_role & UserRole.DEPARTMENT_ADMIN or new_role & UserRole.SUPER_ADMIN or new_role & UserRole.SYSTEM_ADMIN
            
            if was_admin and not is_admin:
                try:
                    departments_with_user = Department.objects.filter(managers=user, is_deleted=False)
                    for dept in departments_with_user:
                        dept.managers.remove(user)
                except Exception:
                    pass
            
            elif old_dept_id and old_dept_id != new_dept_id:
                try:
                    old_department = Department.objects.get(id=old_dept_id, is_deleted=False)
                    old_department.managers.remove(user)
                except Department.DoesNotExist:
                    pass
            
            if is_admin and new_dept_id:
                try:
                    new_department = Department.objects.get(id=new_dept_id, is_deleted=False)
                    if user.id not in new_department.managers.values_list('id', flat=True):
                        new_department.managers.add(user)
                except Department.DoesNotExist:
                    pass

        new_data = {
            'nickname': user.nickname,
            'phone': user.phone,
            'email': user.email,
            'status': user.status,
            'role': user.role,
            'department_id': user.department_id,
        }

        CacheInvalidator.invalidate_user_cache(user_id)

        OperationLogService.log_user_operation(
            request=request,
            operation_type='user_update',
            user=user,
            old_data=old_data,
            new_data=new_data
        )

        return user

    @transaction.atomic
    def delete_user(self, requester, user_id: int, request=None) -> bool:
        try:
            user = User.objects.get(id=user_id, is_deleted=False)
        except User.DoesNotExist:
            raise NotFoundError('用户不存在')
        
        if user.id == requester.id:
            raise ValidationError('不能删除自己')
        
        if not self._can_manage_user(requester, user):
            raise PermissionDenied('无权限删除该用户')

        username = user.nickname or user.username
        user.delete()

        CacheInvalidator.invalidate_user_cache(user_id)

        OperationLogService.log_user_operation(
            request=request,
            operation_type='user_delete',
            user=user,
            description=f'删除了用户 {username}'
        )

        return True

    @transaction.atomic
    def batch_delete_users(self, requester, user_ids: list, request=None) -> dict:
        deleted_count = 0
        failed_list = []
        
        for user_id in user_ids:
            try:
                self.delete_user(requester, user_id)
                deleted_count += 1
            except Exception as e:
                failed_list.append({'id': user_id, 'reason': str(e)})

        if deleted_count > 0:
            CacheInvalidator.invalidate_user_cache(requester.id)
        
        return {
            'deleted_count': deleted_count,
            'failed_count': len(failed_list),
            'failed_list': failed_list,
        }

    @transaction.atomic
    def activate_user(self, requester, user_id: int, is_active: bool, request=None) -> User:
        try:
            user = User.objects.get(id=user_id, is_deleted=False)
        except User.DoesNotExist:
            raise NotFoundError('用户不存在')

        if not self._can_manage_user(requester, user):
            raise PermissionDenied('无权限操作该用户')

        user.is_active = is_active
        user.save(update_fields=['is_active'])

        CacheInvalidator.invalidate_user_cache(user_id)

        OperationLogService.log_user_operation(
            request=request,
            operation_type='user_activate',
            user=user
        )

        return user

    @transaction.atomic
    def update_user_role(self, requester, user_id: int, role: int, department_id: int = None, request=None) -> User:
        if not requester.is_super_admin and not requester.is_department_admin:
            raise PermissionDenied('无权限修改用户角色')

        try:
            user = User.objects.get(id=user_id, is_deleted=False)
        except User.DoesNotExist:
            raise NotFoundError('用户不存在')

        if not self._can_manage_user(requester, user):
            raise PermissionDenied('无权限修改该用户角色')

        old_role = user.role
        old_department_id = user.department_id
        username = user.nickname or user.username
        user.role = role
        if department_id is not None:
            user.department_id = department_id
        user.save(update_fields=['role', 'department_id'])

        CacheInvalidator.invalidate_user_cache(user_id)

        OperationLogService.log_user_operation(
            request=request,
            operation_type='user_role_update',
            user=user,
            description=f'更新了用户 {username} 的权限',
            old_data={'role': old_role, 'department_id': old_department_id},
            new_data={'role': role, 'department_id': user.department_id}
        )

        return user

    @transaction.atomic
    def assign_permission(self, requester, user_id: int, permissions: list) -> User:
        if not requester.is_super_admin:
            raise PermissionDenied('只有超级管理员可以分配权限')
        
        try:
            user = User.objects.get(id=user_id, is_deleted=False)
        except User.DoesNotExist:
            raise NotFoundError('用户不存在')
        
        user.custom_permissions = permissions
        user.save(update_fields=['custom_permissions'])

        CacheInvalidator.invalidate_user_cache(user_id)

        return user

    @transaction.atomic
    def cleanup_user_data(self, requester, user_id: int) -> dict:
        if not requester.is_super_admin:
            raise PermissionDenied('只有超级管理员可以清理用户数据')
        
        try:
            user = User.objects.get(id=user_id, is_deleted=False)
        except User.DoesNotExist:
            raise NotFoundError('用户不存在')
        
        from apps.schedules.models import Schedule
        from apps.records.models import UsageRecord
        from apps.maintenance.models import WorkOrder
        
        cleaned = {
            'schedules': Schedule.objects.filter(teacher=user).delete()[0],
            'records': UsageRecord.objects.filter(teacher=user).delete()[0],
            'work_orders_reporter': WorkOrder.objects.filter(reporter=user).delete()[0],
            'work_orders_handler': WorkOrder.objects.filter(handler=user).update(handler=None),
        }

        CacheInvalidator.invalidate_user_cache(user_id)

        return cleaned

    def get_user_options(self, requester) -> list:
        queryset = User.objects.filter(is_deleted=False, is_active=True)
        
        if requester.is_department_admin and not requester.is_super_admin:
            queryset = queryset.filter(department_id=requester.department_id)
        
        return [
            {
                'id': user.id,
                'username': user.username,
                'nickname': user.nickname or user.username,
                'department_name': user.department.name if user.department else ''
            }
            for user in queryset.order_by('username')
        ]

    @transaction.atomic
    def import_users(self, requester, file_data) -> dict:
        if not requester.is_department_admin and not requester.is_super_admin:
            raise PermissionDenied('无权限导入用户')

        from common.services.progress_service import ProgressService
        
        ProgressService.start_task(
            user_id=requester.id,
            task_type='import_user',
            total=0
        )

        try:
            wb = load_workbook(filename=file_data)
            sheet = wb.active
            
            headers = []
            for cell in sheet[1]:
                headers.append(str(cell.value).strip() if cell.value else '')
            
            required_headers = ['用户名']
            for header in required_headers:
                if header not in headers:
                    ProgressService.fail_task(
                        user_id=requester.id,
                        task_type='import_user',
                        error_message=f'缺少必要列: {header}'
                    )
                    raise ValidationError(f'Excel文件缺少必要列: {header}')

            username_idx = headers.index('用户名')
            nickname_idx = headers.index('昵称') if '昵称' in headers else (headers.index('姓名') if '姓名' in headers else -1)
            phone_idx = headers.index('手机号') if '手机号' in headers else -1
            email_idx = headers.index('邮箱') if '邮箱' in headers else -1
            department_idx = headers.index('所属部门') if '所属部门' in headers else (headers.index('部门') if '部门' in headers else -1)
            role_idx = headers.index('角色') if '角色' in headers else (headers.index('权限') if '权限' in headers else -1)

            rows = list(sheet.iter_rows(min_row=2, values_only=True))
            total_rows = len([r for r in rows if any(r)])
            
            ProgressService.update_progress(
                user_id=requester.id,
                task_type='import_user',
                current=0,
                total=total_rows,
                message='开始导入用户'
            )

            success_count = 0
            failed_count = 0
            failed_list = []
            department_cache = {}

            for idx, row in enumerate(rows):
                if not any(row):
                    continue
                
                try:
                    username = str(row[username_idx]).strip() if row[username_idx] else ''
                    if not username:
                        failed_count += 1
                        failed_list.append({
                            'row': idx + 2,
                            'username': '',
                            'reason': '用户名不能为空'
                        })
                        continue

                    if User.objects.filter(username=username).exists():
                        failed_count += 1
                        failed_list.append({
                            'row': idx + 2,
                            'username': username,
                            'reason': f'用户名 "{username}" 已存在'
                        })
                        continue

                    nickname = str(row[nickname_idx]).strip() if nickname_idx >= 0 and row[nickname_idx] else username
                    phone = str(row[phone_idx]).strip() if phone_idx >= 0 and row[phone_idx] else ''
                    email = str(row[email_idx]).strip() if email_idx >= 0 and row[email_idx] else ''

                    department = None
                    if department_idx >= 0 and row[department_idx]:
                        dept_name = str(row[department_idx]).strip()
                        if dept_name in department_cache:
                            department = department_cache[dept_name]
                        else:
                            try:
                                department = Department.objects.get(name=dept_name)
                                department_cache[dept_name] = department
                            except Department.DoesNotExist:
                                pass
                    
                    if not department and not requester.is_super_admin:
                        department = requester.department

                    if department and not requester.is_super_admin:
                        if department.id != requester.department_id:
                            failed_count += 1
                            failed_list.append({
                                'row': idx + 2,
                                'username': username,
                                'reason': '无权限在该部门创建用户'
                            })
                            continue

                    role = UserRole.TEACHER
                    if role_idx >= 0 and row[role_idx]:
                        role_name = str(row[role_idx]).strip()
                        role_map = {
                            '教师': UserRole.TEACHER,
                            '老师': UserRole.TEACHER,
                            '实训室管理员': UserRole.LABORATORY_ADMIN,
                            '分院管理员': UserRole.DEPARTMENT_ADMIN,
                            '超级管理员': UserRole.SUPER_ADMIN,
                            '系统管理员': UserRole.SYSTEM_ADMIN,
                        }
                        role_names = [r.strip() for r in role_name.replace('、', ' ').replace(',', ' ').split()]
                        role = 0
                        for r in role_names:
                            role |= role_map.get(r, 0)
                        if role == 0:
                            role = UserRole.TEACHER

                    default_password = username[:6] if len(username) >= 6 else username

                    User.objects.create_user(
                        username=username,
                        nickname=nickname,
                        phone=phone,
                        email=email,
                        role=role,
                        department=department,
                        is_active=True,
                        password=default_password
                    )

                    success_count += 1

                    ProgressService.update_progress(
                        user_id=requester.id,
                        task_type='import_user',
                        current=idx + 1,
                        total=total_rows,
                        message=f'正在处理: {username}'
                    )

                except Exception as e:
                    failed_count += 1
                    failed_list.append({
                        'row': idx + 2,
                        'username': str(row[username_idx]) if row[username_idx] else '',
                        'reason': str(e)
                    })

            CacheInvalidator.invalidate_user_cache(requester.id)

            result = {
                'success_count': success_count,
                'failed_count': failed_count,
                'failed_list': failed_list[:100],
                'total': total_rows
            }

            ProgressService.complete_task(
                user_id=requester.id,
                task_type='import_user',
                message=f'导入完成: 成功 {success_count} 个, 失败 {failed_count} 个',
                extra_data=result
            )

            return result

        except Exception as e:
            ProgressService.fail_task(
                user_id=requester.id,
                task_type='import_user',
                error_message=str(e)
            )
            raise

    def _can_manage_user(self, requester, target_user) -> bool:
        if requester.is_super_admin:
            return True
        if requester.is_department_admin:
            return target_user.department_id == requester.department_id
        return False

    def _format_user(self, user) -> dict:
        managed_labs = user.managed_laboratories.filter(is_deleted=False)
        lab_names = [lab.name for lab in managed_labs] if managed_labs.exists() else []
        managed_depts = user.managed_departments.filter(is_deleted=False)
        dept_names = [dept.name for dept in managed_depts] if managed_depts.exists() else []
        return {
            'id': user.id,
            'username': user.username,
            'nickname': user.nickname or user.username,
            'email': user.email or '',
            'phone': user.phone or '',
            'role': user.role,
            'department_id': user.department_id,
            'department_name': user.department.name if user.department else '',
            'managed_departments': '、'.join(dept_names) if dept_names else '暂无管理部门',
            'managed_laboratories': '、'.join(lab_names) if lab_names else '暂无管理实训室',
            'is_active': user.is_active,
            'status': 'active' if user.is_active else 'inactive',
            'created_at': beijing_strftime(user.created_at),
            'last_login': beijing_strftime(user.last_login),
        }
