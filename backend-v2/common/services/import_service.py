"""
导入服务
"""

import io
from django.db import models, transaction
from openpyxl import load_workbook
from apps.core.exceptions import ValidationError, PermissionDenied
from apps.laboratories.models import Laboratory
from apps.schedules.models import Schedule, Semester
from apps.laboratories.models import Equipment
from apps.users.models import User, Department
from common.services.cache_service import CacheInvalidator


class ImportService:
    """导入服务"""

    @staticmethod
    def _parse_excel(file_data):
        """解析 Excel 文件"""
        wb = load_workbook(filename=io.BytesIO(file_data.read()))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

        if len(rows) < 2:
            raise ValidationError(f'文件内容为空或只有表头（实际行数：{len(rows)}）')

        header_row = rows[0]
        header_map = {str(h).strip(): idx for idx, h in enumerate(header_row) if h}

        if not header_map:
            raise ValidationError(f'未识别到有效的表头，第一行内容为：{header_row}')

        data_rows = rows[1:]

        empty_rows = 0
        for i, row in enumerate(data_rows):
            if not any(cell is not None and str(cell).strip() for cell in row):
                empty_rows += 1

        if empty_rows == len(data_rows):
            raise ValidationError(f'所有数据行都是空行，请检查Excel文件内容')

        def get_value(row, *keys):
            for key in keys:
                if key in header_map:
                    val = row[header_map[key]]
                    return str(val).strip() if val is not None else ''
            return ''

        return data_rows, get_value, header_map

    @classmethod
    @transaction.atomic
    def import_laboratories(cls, requester, file_data) -> dict:
        if not requester.is_department_admin and not requester.is_super_admin:
            raise PermissionDenied('无权限导入实训室')

        try:
            data_rows, get_value, header_map = cls._parse_excel(file_data)
        except Exception as e:
            raise ValidationError(f'文件解析失败: {str(e)}')

        success_count = 0
        failed_list = []

        for row_num, row in enumerate(data_rows, start=2):
            try:
                name = get_value(row, '实训室名称', 'name')
                code = get_value(row, '门牌号', 'code')

                if not code or not name:
                    failed_list.append({'row': row_num, 'reason': '门牌号或名称为空'})
                    continue

                existing = Laboratory.objects.filter(code=code).first()
                if existing:
                    if existing.is_deleted:
                        existing.code = f"del_{existing.id}"[:50]
                        existing.save(update_fields=['code'])
                    else:
                        failed_list.append({'row': row_num, 'reason': f'门牌号 "{code}" 已存在'})
                        continue

                department_name = get_value(row, '所属部门', 'department')
                department_id = None
                if department_name:
                    dept = Department.objects.filter(name=department_name).first()
                    if dept:
                        department_id = dept.id

                if not requester.is_super_admin:
                    department_id = requester.department_id

                capacity_str = get_value(row, '工位', 'capacity')
                capacity = int(capacity_str) if capacity_str.isdigit() else 30

                admin_name = get_value(row, '管理员', 'admin', '实训室管理员')
                admin_id = None
                if admin_name:
                    user = User.objects.filter(
                        models.Q(username=admin_name) | models.Q(nickname=admin_name),
                        is_deleted=False
                    ).first()
                    if user:
                        admin_id = user.id

                Laboratory.objects.create(
                    name=name,
                    code=code,
                    capacity=capacity,
                    department_id=department_id,
                    admin_id=admin_id,
                    note=get_value(row, '备注', 'note'),
                )
                success_count += 1

            except Exception as e:
                failed_list.append({'row': row_num, 'reason': str(e)})

        CacheInvalidator.invalidate_laboratory_cache(user_id=requester.id)

        return {
            'success_count': success_count,
            'failed_count': len(failed_list),
            'failed_list': failed_list,
            'imported_count': success_count,
            'skipped_conflicts': 0,
            'errors_count': len(failed_list),
            'detected_headers': list(header_map.keys()),
            'total_rows': len(data_rows),
        }

    @classmethod
    @transaction.atomic
    def import_schedules(cls, requester, file_data, laboratory_id=None) -> dict:
        if not requester.is_department_admin and not requester.is_super_admin and not requester.is_laboratory_admin:
            raise PermissionDenied('无权限导入课表')

        try:
            data_rows, get_value, header_map = cls._parse_excel(file_data)
        except Exception as e:
            raise ValidationError(f'文件解析失败: {str(e)}')

        success_count = 0
        failed_list = []
        conflicts_list = []

        semester = Semester.get_current()
        if not semester:
            raise ValidationError('未设置当前学期')

        for row_num, row in enumerate(data_rows, start=2):
            try:
                course_name = get_value(row, '课程名称', 'course_name')
                laboratory_name = get_value(row, '实训室名称', 'laboratory_name')
                weekday_str = get_value(row, '星期', 'weekday')

                if not course_name or not weekday_str:
                    failed_list.append({'row': row_num, 'reason': '课程名称或星期为空'})
                    continue

                if laboratory_id:
                    try:
                        laboratory = Laboratory.objects.get(id=laboratory_id, is_deleted=False)
                    except Laboratory.DoesNotExist:
                        failed_list.append({'row': row_num, 'reason': '指定的实训室不存在'})
                        continue
                elif laboratory_name:
                    try:
                        laboratory = Laboratory.objects.get(name=laboratory_name, is_deleted=False)
                    except Laboratory.DoesNotExist:
                        failed_list.append({'row': row_num, 'reason': f'实训室 "{laboratory_name}" 不存在'})
                        continue
                else:
                    failed_list.append({'row': row_num, 'reason': '未指定实训室'})
                    continue

                weekday = int(weekday_str) if weekday_str.isdigit() else 1
                time_slot = get_value(row, '节次', 'time_slot') or '1-4'
                weeks = get_value(row, '周次', 'weeks') or '1-18'

                existing = Schedule.objects.filter(
                    laboratory=laboratory,
                    semester=semester,
                    course_name=course_name,
                    weekday=weekday,
                    time_slot=time_slot,
                    weeks=weeks,
                    is_deleted=False
                ).first()

                if existing:
                    conflicts_list.append({
                        'row': row_num,
                        'course_name': course_name,
                        'laboratory': laboratory.name,
                        'weekday': weekday,
                        'time_slot': time_slot,
                        'weeks': weeks,
                        'existing_id': existing.id
                    })
                    continue

                Schedule.objects.create(
                    course_name=course_name,
                    weekday=weekday,
                    time_slot=time_slot,
                    weeks=weeks,
                    laboratory=laboratory,
                    semester=semester,
                    teacher_name=get_value(row, '任课教师', 'teacher_name'),
                    class_name=get_value(row, '上课班级', 'class_name'),
                    student_count=int(get_value(row, '人数', 'student_count') or 0),
                    note=get_value(row, '备注', 'note'),
                )
                success_count += 1

            except Exception as e:
                failed_list.append({'row': row_num, 'reason': str(e)})

        return {
            'success_count': success_count,
            'failed_count': len(failed_list),
            'failed_list': failed_list,
            'imported_count': success_count,
            'skipped_conflicts': len(conflicts_list),
            'conflicts_list': conflicts_list,
            'errors_count': len(failed_list),
            'detected_headers': list(header_map.keys()),
            'total_rows': len(data_rows),
            'required_fields': ['实训室名称', '课程名称', '星期'],
            'hint': '请确保Excel表头与系统要求的列名一致，实训室必须已存在于系统中',
        }

    @classmethod
    @transaction.atomic
    def import_equipment(cls, requester, file_data) -> dict:
        if not requester.is_department_admin and not requester.is_super_admin and not requester.is_laboratory_admin:
            raise PermissionDenied('无权限导入设备')

        try:
            data_rows, get_value, header_map = cls._parse_excel(file_data)
        except Exception as e:
            raise ValidationError(f'文件解析失败: {str(e)}')

        success_count = 0
        failed_list = []

        for row_num, row in enumerate(data_rows, start=2):
            try:
                code = get_value(row, '电脑编号', 'code')
                name = get_value(row, '设备名称', 'name')

                if not code:
                    failed_list.append({'row': row_num, 'reason': '电脑编号为空'})
                    continue

                laboratory_name = get_value(row, '实训室名称', 'laboratory_name')
                laboratory_id = None
                if laboratory_name:
                    try:
                        lab = Laboratory.objects.get(name=laboratory_name, is_deleted=False)
                        laboratory_id = lab.id
                    except Laboratory.DoesNotExist:
                        failed_list.append({'row': row_num, 'reason': f'实训室 "{laboratory_name}" 不存在'})
                        continue

                status_value = get_value(row, '状态', 'status')
                if status_value:
                    status_map = {
                        '正常': 'NORMAL',
                        '故障': 'DAMAGED',
                        '维修中': 'MAINTENANCE',
                        '维护中': 'MAINTENANCE',
                        '已报废': 'SCRAPPED',
                        '借用中': 'BORROWED',
                    }
                    status_value = status_map.get(status_value, status_value)

                Equipment.objects.create(
                    name=name or code,
                    code=code,
                    category=get_value(row, '类型', 'category') or '计算机',
                    brand=get_value(row, '品牌', 'brand'),
                    model=get_value(row, '型号', 'model'),
                    laboratory_id=laboratory_id,
                    cpu=get_value(row, 'CPU', 'cpu'),
                    memory=get_value(row, '内存', 'memory'),
                    disk=get_value(row, '硬盘', 'disk'),
                    status=status_value or 'NORMAL',
                    note=get_value(row, '备注', 'note'),
                )
                success_count += 1

            except Exception as e:
                failed_list.append({'row': row_num, 'reason': str(e)})

        CacheInvalidator.invalidate_laboratory_cache(user_id=requester.id)

        return {
            'success_count': success_count,
            'failed_count': len(failed_list),
            'failed_list': failed_list,
            'imported_count': success_count,
            'skipped_conflicts': 0,
            'errors_count': len(failed_list),
            'detected_headers': list(header_map.keys()),
            'total_rows': len(data_rows),
        }
